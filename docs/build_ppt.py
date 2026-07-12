# -*- coding: utf-8 -*-
"""
生成《SmartCode Browser 项目介绍》PPT。

叙事主线（本版的组织逻辑）：
  缝隙（痛点）→ 级联交互填它（demo）→ 用 AI「聊」出来的（方法论·重点）
  → 踩过的坑（注意事项）→ 这不是玩具（设计取舍）→ 往哪走（路线图）

取向：这是个几乎全程由 AI(Cursor) 写出来的工具，「代码实现」不是重点；
重点是「为什么需要、怎么造的、要注意什么、往哪走」，集中在最值得分享的部分。

运行：.venv/bin/python docs/build_ppt.py
产物：docs/SmartCode_Browser_介绍.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ---------------------------------------------------------------------------
# 主题：深色现代风，主色取自工具里「可点击调用」的蓝
# ---------------------------------------------------------------------------
BG_DARK = RGBColor(0x0E, 0x16, 0x27)
BG_PANEL = RGBColor(0x17, 0x22, 0x38)
BG_PANEL2 = RGBColor(0x1F, 0x2D, 0x47)
ACCENT = RGBColor(0x4D, 0x9C, 0xFF)
ACCENT2 = RGBColor(0x32, 0xD6, 0xC0)
WARN = RGBColor(0xFF, 0xB1, 0x4D)
TEXT = RGBColor(0xEC, 0xF1, 0xF8)
MUTED = RGBColor(0x93, 0xA1, 0xB5)
LINE = RGBColor(0x2A, 0x3A, 0x57)

CJK = "Microsoft YaHei"
MONO = "Consolas"

SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)
IN = 914400

OUT = Path(__file__).resolve().parent / "SmartCode_Browser_介绍.pptx"


# ---------------------------------------------------------------------------
# 底层工具
# ---------------------------------------------------------------------------
def _set_font(run, size=18, color=TEXT, bold=False, italic=False, font=CJK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    rect.fill.solid()
    rect.fill.fore_color.rgb = BG_DARK
    rect.line.fill.background()
    rect.shadow.inherit = False
    sp = rect._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)
    return slide


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=True):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(st, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    return tf


def para(tf, text, size=18, color=TEXT, bold=False, italic=False, align=PP_ALIGN.LEFT,
         font=CJK, space_after=8, space_before=0, first=False, line_spacing=None):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if space_after is not None:
        p.space_after = Pt(space_after)
    if space_before:
        p.space_before = Pt(space_before)
    if line_spacing:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    _set_font(r, size=size, color=color, bold=bold, italic=italic, font=font)
    return p


def runs(p, segments, align=PP_ALIGN.LEFT, space_after=8):
    p.alignment = align
    p.space_after = Pt(space_after)
    for text, style in segments:
        r = p.add_run()
        r.text = text
        _set_font(r, **style)


def header(slide, kicker, title, accent=ACCENT):
    box(slide, int(0.7 * IN), int(0.62 * IN), int(0.10 * IN), int(0.55 * IN), fill=accent, radius=False)
    tf = textbox(slide, int(0.95 * IN), int(0.5 * IN), int(11.6 * IN), int(1.3 * IN))
    para(tf, kicker, size=14, color=accent, bold=True, first=True, space_after=2)
    para(tf, title, size=29, color=TEXT, bold=True, space_after=0)


def footer(slide, idx, total):
    tf = textbox(slide, int(0.7 * IN), int(6.95 * IN), int(11.0 * IN), int(0.35 * IN))
    runs(tf.paragraphs[0],
         [("SmartCode Browser", dict(size=10, color=MUTED, bold=True)),
          ("   ·   一个用 Cursor「聊」出来的代码导航工具", dict(size=10, color=MUTED))],
         space_after=0)
    tf2 = textbox(slide, int(11.2 * IN), int(6.95 * IN), int(1.3 * IN), int(0.35 * IN))
    para(tf2, f"{idx:02d} / {total:02d}", size=10, color=MUTED, align=PP_ALIGN.RIGHT, first=True, space_after=0)


def bullet(tf, text, size=18, color=TEXT, bold=False, dot="·", dot_color=ACCENT,
           space_after=10, first=False):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = 1.05
    r0 = p.add_run(); r0.text = dot + "  "
    _set_font(r0, size=size, color=dot_color, bold=True)
    r1 = p.add_run(); r1.text = text
    _set_font(r1, size=size, color=color, bold=bold)
    return p


def card(slide, x, y, w, h, title, lines, accent=ACCENT, title_size=18, body_size=14):
    box(slide, x, y, w, h, fill=BG_PANEL, line=LINE, line_w=1.0)
    box(slide, x, y, int(0.09 * IN), h, fill=accent, radius=False)
    tf = textbox(slide, x + int(0.28 * IN), y + int(0.16 * IN), w - int(0.5 * IN), h - int(0.3 * IN))
    para(tf, title, size=title_size, color=TEXT, bold=True, first=True, space_after=8)
    for ln in lines:
        bullet(tf, ln, size=body_size, color=MUTED, dot="·", dot_color=accent, space_after=6)


def simple_table(slide, x, y, w, headers, rows, col_ratios, row_h=0.5, header_h=0.55,
                 head_fill=BG_PANEL2, body_size=14, head_size=15):
    total = sum(col_ratios)
    col_w = [int(w * r / total) for r in col_ratios]
    cx = x
    for j, ht in enumerate(headers):
        c = box(slide, cx, y, col_w[j], int(header_h * IN), fill=head_fill, line=BG_DARK, line_w=1.5)
        tf = c.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, ht, size=head_size, color=ACCENT if j == 0 else TEXT, bold=True, first=True, space_after=0)
        tf.margin_left = Pt(10)
        cx += col_w[j]
    ry = y + int(header_h * IN)
    for i, row in enumerate(rows):
        cx = x
        rfill = BG_PANEL if i % 2 == 0 else BG_PANEL2
        for j, cell in enumerate(row):
            c = box(slide, cx, ry, col_w[j], int(row_h * IN), fill=rfill, line=BG_DARK, line_w=1.5)
            tf = c.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Pt(10); tf.margin_right = Pt(8)
            para(tf, cell, size=body_size, color=TEXT if j == 0 else MUTED, bold=(j == 0),
                 first=True, space_after=0, line_spacing=1.0)
            cx += col_w[j]
        ry += int(row_h * IN)
    return ry


def arrow(slide, x, y, w, h, color=ACCENT2):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(x), Emu(y), Emu(w), Emu(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a


def chev(slide, x, y, w, h, color=MUTED):
    a = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Emu(x), Emu(y), Emu(w), Emu(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a


def stat(slide, x, y, w, number, label, accent=ACCENT):
    box(slide, x, y, w, int(1.25 * IN), fill=BG_PANEL, line=LINE)
    tf = textbox(slide, x, y + int(0.16 * IN), w, int(1.0 * IN), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, number, size=30, color=accent, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=2)
    para(tf, label, size=12, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)


def quad(slide, x, y, w, h, idx, title, lines, accent):
    """注意事项页用的象限卡：序号 + 标题 + 要点。"""
    box(slide, x, y, w, h, fill=BG_PANEL, line=LINE)
    box(slide, x, y, int(0.09 * IN), h, fill=accent, radius=False)
    tfh = textbox(slide, x + int(0.28 * IN), y + int(0.14 * IN), w - int(0.5 * IN), int(0.5 * IN))
    runs(tfh.paragraphs[0],
         [(idx + "  ", dict(size=18, color=accent, bold=True)),
          (title, dict(size=17, color=TEXT, bold=True))], space_after=0)
    tfb = textbox(slide, x + int(0.28 * IN), y + int(0.66 * IN), w - int(0.5 * IN), h - int(0.78 * IN))
    for k, ln in enumerate(lines):
        bullet(tfb, ln, size=13, color=MUTED, dot="·", dot_color=accent, space_after=6, first=(k == 0))


# ===========================================================================
prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H


def n():
    return len(prs.slides._sldIdLst)


# ---- 1. 封面 -------------------------------------------------------------
s = add_slide(prs)
box(s, 0, 0, SLIDE_W, int(0.18 * IN), fill=ACCENT, radius=False)
box(s, 0, int(0.18 * IN), SLIDE_W, int(0.06 * IN), fill=ACCENT2, radius=False)
tf = textbox(s, int(0.9 * IN), int(1.8 * IN), int(11.2 * IN), int(2.9 * IN))
para(tf, "SmartCode  Browser", size=52, color=TEXT, bold=True, first=True, space_after=6)
para(tf, "级联式代码浏览器", size=29, color=ACCENT, bold=True, space_after=18)
para(tf, "一个人 + AI，做出一个真能用在内核 / 芯片大库上的代码导航工具", size=19, color=TEXT, space_after=4)
para(tf, "几乎没手写代码 —— 全程用 Cursor「聊」出来", size=16, color=MUTED, space_after=0)
box(s, int(0.9 * IN), int(5.4 * IN), int(10.6 * IN), 0, line=LINE, line_w=1.0, radius=False)
tf2 = textbox(s, int(0.9 * IN), int(5.6 * IN), int(11 * IN), int(1.0 * IN))
runs(tf2.paragraphs[0],
     [("为什么需要 ", dict(size=16, color=ACCENT2, bold=True)),
      ("·  怎么用 AI 造的  ·  ", dict(size=16, color=MUTED)),
      ("要注意什么 ", dict(size=16, color=ACCENT2, bold=True)),
      ("·  往哪走", dict(size=16, color=MUTED))], space_after=0)

# ---- 2. 主线（story spine）----------------------------------------------
s = add_slide(prs)
header(s, "今天的主线", "一条线串起来：从缝隙到工具，再到 AI 怎么造的")
spine = [
    ("缝隙", "现有工具追一条链都别扭", ACCENT),
    ("级联", "用新交互填上，现场演示", ACCENT),
    ("AI 造的", "全程和 Cursor 聊出来", WARN),
    ("踩过的坑", "用 AI 开发要注意什么", WARN),
    ("往哪走", "路线图", ACCENT2),
]
bw = int(2.05 * IN); gap = int(0.28 * IN); x = int(0.95 * IN); y = int(2.4 * IN); bh = int(1.7 * IN)
for i, (t, d, ac) in enumerate(spine):
    box(s, x, y, bw, bh, fill=BG_PANEL, line=ac, line_w=1.5)
    box(s, x, y, bw, int(0.12 * IN), fill=ac, radius=False)
    tf = textbox(s, x + int(0.12 * IN), y + int(0.2 * IN), bw - int(0.24 * IN), bh - int(0.35 * IN),
                 anchor=MSO_ANCHOR.MIDDLE)
    para(tf, f"{i+1}", size=14, color=ac, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=4)
    para(tf, t, size=20, color=TEXT, bold=True, align=PP_ALIGN.CENTER, space_after=6)
    para(tf, d, size=12, color=MUTED, align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.05)
    if i < len(spine) - 1:
        chev(s, x + bw - int(0.02 * IN), y + bh // 2 - int(0.18 * IN), int(0.34 * IN), int(0.36 * IN),
             color=MUTED)
    x += bw + gap
tf = textbox(s, int(0.95 * IN), int(4.6 * IN), int(11.3 * IN), int(1.0 * IN))
runs(tf.paragraphs[0],
     [("重点在第 3 段：", dict(size=16, color=WARN, bold=True)),
      ("这工具本身只是个例子，更想分享的是 —— 现在你也能这样用 AI 造出真能用的内部工具。",
       dict(size=16, color=TEXT))], space_after=0)

# ===========================================================================
# A · 为什么 + 是什么（钩子 + demo）
# ===========================================================================

# ---- 3. 为什么需要它：缝隙洞察（钩子）------------------------------------
s = add_slide(prs)
header(s, "为什么需要它", "「调查一条调用链」是被现有工具组织得最烂的任务")
simple_table(
    s, int(0.95 * IN), int(1.7 * IN), int(10.6 * IN),
    ["现有方式", "能干嘛", "追一条链时的别扭"],
    [
        ["IDE 跳转", "准、能改代码", "标签越开越多，路径全在脑子里，难复述给别人"],
        ["grep / 全文搜索", "快、无需索引", "只给「出现位置」，没有调用层次"],
        ["直接问 AI", "出解释快", "行号 / 路径不稳定，难逐行核对"],
    ],
    col_ratios=[2.4, 2.8, 5.4], row_h=0.78, header_h=0.6, body_size=15,
)
box(s, int(0.95 * IN), int(4.85 * IN), int(10.6 * IN), int(1.5 * IN), fill=BG_PANEL, line=ACCENT, line_w=1.5)
tf = textbox(s, int(1.3 * IN), int(5.0 * IN), int(10 * IN), int(1.2 * IN), anchor=MSO_ANCHOR.MIDDLE)
para(tf, "缝隙：不是工具不好，而是「沿调用链调查」这件事，信息组织方式不对。",
     size=18, color=TEXT, bold=True, first=True, space_after=6)
runs(tf.add_paragraph(),
     [("我们想要的：", dict(size=15, color=ACCENT, bold=True)),
      ("从入口出发，顺着「谁调谁」一层层展开，路径可见、可续、可分享 —— 像在树上散步。",
       dict(size=15, color=MUTED))], space_after=0)

# ---- 4. 核心体验：级联交互（demo 的 aha）--------------------------------
s = add_slide(prs)
header(s, "怎么填这个缝隙", "核心体验：点蓝色调用，右侧长出新面板")
py = int(2.1 * IN); ph = int(2.6 * IN); pw = int(3.2 * IN)
labels = [("函数 A", "入口 / 搜索进入", ACCENT, True),
          ("函数 B", "点 A 里的调用打开", ACCENT, True),
          ("函数 C", "再点 B 里的调用打开", ACCENT2, False)]
xs = [int(0.95 * IN), int(4.55 * IN), int(8.15 * IN)]
for i, (t, d, ac, has) in enumerate(labels):
    box(s, xs[i], py, pw, ph, fill=BG_PANEL, line=ac, line_w=1.5)
    box(s, xs[i], py, pw, int(0.55 * IN), fill=ac, radius=True)
    box(s, xs[i], py + int(0.27 * IN), pw, int(0.28 * IN), fill=ac, radius=False)
    tfh = textbox(s, xs[i], py + int(0.05 * IN), pw, int(0.5 * IN), anchor=MSO_ANCHOR.MIDDLE)
    para(tfh, t, size=18, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=0)
    tfb = textbox(s, xs[i] + int(0.25 * IN), py + int(0.75 * IN), pw - int(0.5 * IN), int(1.7 * IN))
    para(tfb, d, size=13, color=MUTED, first=True, space_after=10)
    if has:
        runs(tfb.add_paragraph(),
             [("  result = ", dict(size=13, color=TEXT, font=MONO)),
              ("call_next()", dict(size=13, color=ACCENT, bold=True, font=MONO)),
              (";", dict(size=13, color=TEXT, font=MONO))], space_after=0)
for i in range(2):
    arrow(s, xs[i] + pw + int(0.02 * IN), py + int(1.0 * IN), int(0.55 * IN), int(0.5 * IN), color=ACCENT2)
tf = textbox(s, int(0.95 * IN), int(5.1 * IN), int(11 * IN), int(1.2 * IN))
runs(tf.paragraphs[0],
     [("新面板向右生长，左侧不被盖住", dict(size=17, color=ACCENT, bold=True)),
      (" —— 整条调用链 A→B→C 始终可见、可回看、可截图分享。", dict(size=17, color=TEXT))], space_after=0)

# ---- 5. 能力一览（压缩，一页）-------------------------------------------
s = add_slide(prs)
header(s, "它能做什么", "能力一览")
simple_table(
    s, int(0.95 * IN), int(1.7 * IN), int(10.6 * IN),
    ["能力", "一句话说明"],
    [
        ["调用 → 定义", "点蓝色引用，级联展开；多实现时真实现优先、空桩靠后"],
        ["变量 → 声明", "局部变量本面板高亮，全局 / 字段 / 枚举开新面板"],
        ["查找用法", "Alt/Ctrl+点击 或右键，看谁在用这个符号"],
        ["符号搜索", "支持 in:路径 限定范围，从任意入口开始追链"],
        ["多 arch 消歧", "配 compile_commands.json 后跳到真正被编译的那一份"],
        ["会话 / 标签", "刷新自动恢复 + 命名保存探索路径，方便交接分享"],
        ["可选 AI 分析", "对当前函数提问，走 cursor-agent（复用 Cursor 登录）"],
    ],
    col_ratios=[3.0, 7.6], row_h=0.6, header_h=0.55, body_size=14, head_size=15,
)
para(textbox(s, int(0.95 * IN), int(6.3 * IN), int(11 * IN), int(0.5 * IN)),
     "语言无关：C/C++ · Python · Java · Scala/Chisel —— 换语言、换项目主要靠配置。",
     size=14, color=MUTED, first=True, space_after=0)

# ---- 6. 现场演示（含吃狗粮）---------------------------------------------
s = add_slide(prs)
header(s, "现场演示", "三分钟看明白（重点两个动作）", accent=ACCENT2)
steps = [
    ("1", "搜入口 → 连点 2~3 层调用", "面板一块块向右长出来 —— 全场最直观的瞬间，口播调用链"),
    ("2", "保存标签 → 刷新恢复", "命名保存路径 → F5，阅读树自动恢复，强调「可续」"),
    ("★", "吃自己的狗粮", "当场用这个工具去读它自己的源码 —— AI 做的工具来读代码，最有记忆点"),
]
y = int(2.0 * IN)
for num, t, d in steps:
    star = num == "★"
    ac = WARN if star else ACCENT2
    box(s, int(0.95 * IN), y, int(0.7 * IN), int(1.05 * IN), fill=ac)
    tfn = s.shapes[-1].text_frame; tfn.word_wrap = True; tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tfn, num, size=26, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=0)
    box(s, int(1.85 * IN), y, int(9.7 * IN), int(1.05 * IN), fill=BG_PANEL, line=LINE)
    if star:
        box(s, int(1.85 * IN), y, int(0.09 * IN), int(1.05 * IN), fill=WARN, radius=False)
    tf = textbox(s, int(2.15 * IN), y + int(0.16 * IN), int(9.2 * IN), int(0.8 * IN), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, t, size=18, color=TEXT, bold=True, first=True, space_after=3)
    para(tf, d, size=13.5, color=MUTED, space_after=0)
    y += int(1.27 * IN)
para(textbox(s, int(0.95 * IN), int(6.05 * IN), int(11 * IN), int(0.5 * IN)),
     "其它（查用法、多候选选真实现、AI 提问）视时间补充，别让演示掉节奏。",
     size=13, color=MUTED, italic=True, first=True, space_after=0)

# ===========================================================================
# B · 怎么用 AI 造出来的（重点，加重）
# ===========================================================================

# ---- 7. 分隔：用 Cursor 聊出来的 + 数字 ---------------------------------
s = add_slide(prs)
box(s, 0, int(1.9 * IN), SLIDE_W, int(1.7 * IN), fill=BG_PANEL, radius=False)
box(s, int(0.9 * IN), int(1.9 * IN), int(0.14 * IN), int(1.7 * IN), fill=WARN, radius=False)
tf = textbox(s, int(1.3 * IN), int(2.0 * IN), int(10.5 * IN), int(1.5 * IN), anchor=MSO_ANCHOR.MIDDLE)
para(tf, "PART · 重点", size=16, color=WARN, bold=True, first=True, space_after=6)
para(tf, "怎么用 Cursor 把这个工具「聊」出来", size=33, color=TEXT, bold=True, space_after=0)
stats = [("≈6900", "行代码"), ("8", "个 Python 模块"), ("~2400", "行前端"), ("7+", "次能力迭代")]
sw = int(2.55 * IN); x0 = int(0.95 * IN); sy = int(4.1 * IN)
for i, (num, lab) in enumerate(stats):
    stat(s, x0 + i * (sw + int(0.2 * IN)), sy, sw, num, lab, accent=WARN if i % 2 else ACCENT)
para(textbox(s, int(0.95 * IN), int(5.6 * IN), int(11 * IN), int(0.5 * IN)),
     "几乎全程 AI 结对完成 —— 人几乎没手写实现，只在「说清楚要什么」和「把关」上花时间。",
     size=14, color=MUTED, first=True, space_after=0)

# ---- 8. 核心理念 ---------------------------------------------------------
s = add_slide(prs)
header(s, "怎么造的 (1/4)", "核心理念：不写实现，写「想要什么」", accent=WARN)
cards2 = [
    ("从真实痛点出发", ["起点是「读 cpuidle 调用链很累」", "不是「我想学个新框架」", "需求来自真实工程，验收也回真实工程"], WARN),
    ("用自然语言描述需求", ["「点调用就在右边展开定义」", "「同名多实现时真实现优先」", "AI 负责翻译成跨文件的代码"], ACCENT),
    ("人审 + 小步迭代", ["每个能力一个 commit，独立验证", "跑起来不对就回退、重新描述", "实现是 AI 的活儿，方向是人的活儿"], ACCENT2),
]
cw = int(3.6 * IN); x0 = int(0.95 * IN); cy = int(1.9 * IN); ch = int(3.4 * IN)
for i, (t, lines, ac) in enumerate(cards2):
    card(s, x0 + i * (cw + int(0.25 * IN)), cy, cw, ch, t, lines, accent=ac, title_size=18, body_size=15)
tf = textbox(s, int(0.95 * IN), int(5.6 * IN), int(11 * IN), int(0.7 * IN))
runs(tf.paragraphs[0],
     [("一句话：", dict(size=16, color=WARN, bold=True)),
      ("把 Cursor 当「随叫随到的结对工程师」—— 你负责想清楚，它负责写出来。", dict(size=16, color=TEXT))],
     space_after=0)

# ---- 9. 真实迭代时间线 ---------------------------------------------------
s = add_slide(prs)
header(s, "怎么造的 (2/4)", "真实迭代时间线（每条 = 一次对话 + 一个 commit）", accent=WARN)
timeline = [
    ("v0", "独立级联浏览器", "先把「点调用→右侧展开」的骨架跑起来", ACCENT),
    ("+", "面板去重 · 拖动避让 · 会话标签", "让多面板好用：刷新恢复、命名保存路径", ACCENT),
    ("+", "AI 分析侧栏（cursor-agent）", "把 Cursor 接进浏览器，对当前函数提问", ACCENT2),
    ("+", "搜索支持 in:路径 限定", "大库里更快定位入口符号", ACCENT2),
    ("+", "函数指针跳转 · 宏修复 · 索引缓存", "真用内核时撞到的边界，逐个补上 + 提速", WARN),
    ("+", "查用法 · 连线锚点优化", "往回看依赖 + 视觉打磨", WARN),
    ("now", "多入口跳同一函数时复用面板", "细节体验持续抛光", ACCENT),
]
y = int(1.8 * IN); lh = int(0.66 * IN)
box(s, int(1.5 * IN), y, int(0.03 * IN), lh * len(timeline) - int(0.2 * IN), fill=LINE, radius=False)
for tag, t, d, ac in timeline:
    box(s, int(1.28 * IN), y + int(0.06 * IN), int(0.46 * IN), int(0.46 * IN), fill=ac)
    tfn = s.shapes[-1].text_frame; tfn.word_wrap = True; tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tfn, tag, size=12, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=0)
    tf = textbox(s, int(2.0 * IN), y, int(9.5 * IN), lh, anchor=MSO_ANCHOR.MIDDLE)
    runs(tf.paragraphs[0],
         [(t + "    ", dict(size=16, color=TEXT, bold=True)), (d, dict(size=13, color=MUTED))], space_after=0)
    y += lh

# ---- 10. 怎么和 Cursor 配合 ---------------------------------------------
s = add_slide(prs)
header(s, "怎么造的 (3/4)", "具体怎么和 Cursor 配合", accent=WARN)
ways = [
    ("Agent / Composer 多文件改", "一句需求同时改后端 + 前端，保持端到端一致"),
    ("@ 引用精确喂上下文", "@文件 / @符号 指给 AI，少绕弯、改得准"),
    ("边写边跑自检", "curl 接口 + 浏览器实时看效果，错了立刻反馈给 AI"),
    ("SpecStory 留存对话", ".specstory/ 自动记录每次会话 = 开发决策日志"),
    ("吃自己的狗粮", "registry 默认指向本仓库，用工具读工具自身代码"),
    ("让 AI 同步写文档", "README、这份 PPT 讲稿都让 AI 跟着代码一起更新"),
]
cw = int(5.2 * IN); ch = int(1.35 * IN); x0 = int(0.95 * IN); y0 = int(1.85 * IN)
for i, (t, d) in enumerate(ways):
    col = i % 2; row = i // 2
    x = x0 + col * (cw + int(0.2 * IN)); y = y0 + row * (ch + int(0.18 * IN))
    box(s, x, y, cw, ch, fill=BG_PANEL, line=LINE)
    box(s, x, y, int(0.09 * IN), ch, fill=WARN if col else ACCENT, radius=False)
    tf = textbox(s, x + int(0.28 * IN), y + int(0.12 * IN), cw - int(0.5 * IN), ch - int(0.24 * IN),
                 anchor=MSO_ANCHOR.MIDDLE)
    para(tf, t, size=16, color=TEXT, bold=True, first=True, space_after=4)
    para(tf, d, size=13, color=MUTED, space_after=0)

# ---- 11. 人 vs AI 分工 ---------------------------------------------------
s = add_slide(prs)
header(s, "怎么造的 (4/4)", "人和 AI 各干什么", accent=WARN)
box(s, int(0.95 * IN), int(1.9 * IN), int(5.15 * IN), int(3.7 * IN), fill=BG_PANEL, line=ACCENT, line_w=1.5)
tfa = textbox(s, int(1.25 * IN), int(2.1 * IN), int(4.6 * IN), int(3.3 * IN))
para(tfa, "人负责（少而关键）", size=20, color=ACCENT, bold=True, first=True, space_after=10)
for t in ["定方向：解决哪个真实痛点", "定边界：只读 / 语言无关 / 不做 IDE",
          "定优先级：如「真实现优先于空桩」", "验真场景：拿内核 / NEMU 实测", "审结果、把关取舍与质量"]:
    bullet(tfa, t, size=15, color=TEXT, dot="·", dot_color=ACCENT, space_after=9)
box(s, int(6.4 * IN), int(1.9 * IN), int(5.15 * IN), int(3.7 * IN), fill=BG_PANEL, line=ACCENT2, line_w=1.5)
tfb = textbox(s, int(6.7 * IN), int(2.1 * IN), int(4.6 * IN), int(3.3 * IN))
para(tfb, "AI（Cursor）负责（重活）", size=20, color=ACCENT2, bold=True, first=True, space_after=10)
for t in ["几乎全部实现：跨文件写 / 改代码", "脚手架：服务、前端、语言适配", "补盲区：宏、函数指针等边界 case",
          "重构、命名、写注释 / 文档", "解释报错、给修复方案"]:
    bullet(tfb, t, size=15, color=TEXT, dot="·", dot_color=ACCENT2, space_after=9)
tf = textbox(s, int(0.95 * IN), int(5.85 * IN), int(11 * IN), int(0.6 * IN))
para(tf, "AI 写了几乎所有代码，但「这东西该是什么样」始终是人在拿主意。",
     size=16, color=TEXT, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=0)

# ---- 12. 用 AI 开发的注意事项（新增，重点）------------------------------
s = add_slide(prs)
header(s, "踩过的坑", "用 AI 做这类东西，要注意什么", accent=WARN)
quads = [
    ("①", "方向与边界：人来守", ["场景要具体（需求 ≈ 验收标准）", "边界钉死，否则范围一路膨胀", "优先级是人的判断，AI 给不了"], ACCENT),
    ("②", "质量：别信「看起来对」", ["每次都审 diff，尤其跨文件", "小步提交 + 拿真实数据验证", "防幻觉：编造的 API 去查文档"], WARN),
    ("③", "上下文：喂得准、会漂", ["@文件 / @符号 精确给上下文", "长对话会忘约定，适时开新会话", "把边界 / 风格固化进规则文件"], ACCENT2),
    ("④", "安全与合规", ["别把密钥 / 内部源码无意识喂出去", "生成代码可能撞开源，关键处留意", "危险操作（删库 / 强推）人来确认"], WARN),
]
cw = int(5.2 * IN); ch = int(2.0 * IN); x0 = int(0.95 * IN); y0 = int(1.8 * IN)
for i, (idx, t, lines, ac) in enumerate(quads):
    col = i % 2; row = i // 2
    x = x0 + col * (cw + int(0.2 * IN)); y = y0 + row * (ch + int(0.18 * IN))
    quad(s, x, y, cw, ch, idx, t, lines, ac)
para(textbox(s, int(0.95 * IN), int(6.0 * IN), int(11 * IN), int(0.5 * IN)),
     "一句话：AI 是放大器 —— 放大你的判断力，也放大你的疏忽。坦诚讲局限，反而最加分。",
     size=14, color=MUTED, italic=True, first=True, space_after=0)

# ---- 13. 这不是玩具：设计取舍（建立可信度）------------------------------
s = add_slide(prs)
header(s, "这不是玩具", "几个为真实工程做的取舍（AI 实现，人来定）")
tradeoffs = [
    ("多 arch / 多 ISA 消歧", "同名实现 x86 / riscv / mips 各一份；读 compile_commands.json 判断「这次实际编译了哪份」，稳定命中真实现", ACCENT),
    ("真实现优先、空桩靠后", "大库里同名 / 弱符号 / #ifdef 桩满天飞 —— 候选排序体现对真实工程的理解", ACCENT2),
    ("诚实优于炫技", "静态解析有盲区，解析不了就明确标注、绝不假装能跳；grep 兜底保证至少搜得到名字", WARN),
]
y = int(1.85 * IN)
for t, d, ac in tradeoffs:
    box(s, int(0.95 * IN), y, int(10.6 * IN), int(1.4 * IN), fill=BG_PANEL, line=LINE)
    box(s, int(0.95 * IN), y, int(0.1 * IN), int(1.4 * IN), fill=ac, radius=False)
    tf = textbox(s, int(1.35 * IN), y + int(0.18 * IN), int(10 * IN), int(1.1 * IN), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, t, size=18, color=ac, bold=True, first=True, space_after=5)
    para(tf, d, size=14, color=MUTED, space_after=0, line_spacing=1.05)
    y += int(1.55 * IN)

# ===========================================================================
# C · 往哪走 + 收尾
# ===========================================================================

# ---- 14. 路线图 ----------------------------------------------------------
s = add_slide(prs)
header(s, "往哪走 · ROADMAP", "后面打算做什么", accent=ACCENT2)
roadmap = [
    ("近期 · Now/Next", ACCENT,
     ["难解析场景增强：函数指针 / 表驱动 / 宏 跳得更准",
      "全局调用图视图：不止级联，整张图一眼看全",
      "调用树导出 / 分享链接：把路径发给同事"]),
    ("中期", ACCENT2,
     ["更深 AI：一键讲解整条调用链",
      "对调用图提问：「谁最终调用了它」",
      "接 IDE / Git / PR：跳编辑器、评审挂调用树"]),
    ("远期", WARN,
     ["更多语言：Rust / Go / Verilog 完善",
      "团队协作：登录鉴权 / 共享标签",
      "性能：超大库增量索引、更快冷启动"]),
]
cw = int(3.6 * IN); x0 = int(0.95 * IN); cy = int(1.85 * IN); ch = int(4.3 * IN)
for i, (t, ac, lines) in enumerate(roadmap):
    x = x0 + i * (cw + int(0.25 * IN))
    box(s, x, cy, cw, ch, fill=BG_PANEL, line=LINE)
    box(s, x, cy, cw, int(0.7 * IN), fill=ac, radius=True)
    box(s, x, cy + int(0.35 * IN), cw, int(0.35 * IN), fill=ac, radius=False)
    tfh = textbox(s, x, cy + int(0.08 * IN), cw, int(0.6 * IN), anchor=MSO_ANCHOR.MIDDLE)
    para(tfh, t, size=18, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=0)
    tfb = textbox(s, x + int(0.3 * IN), cy + int(0.95 * IN), cw - int(0.6 * IN), int(3.2 * IN))
    for k, ln in enumerate(lines):
        bullet(tfb, ln, size=14.5, color=TEXT, dot="▸", dot_color=ac, space_after=14, first=(k == 0))

# ---- 15. 总结（主线回收）------------------------------------------------
s = add_slide(prs)
header(s, "总结", "一页带走")
rows3 = [
    ("缝隙 → 工具", "「沿调用链调查」被现有工具组织得很烂；级联交互把它变成沿调用树散步", ACCENT),
    ("怎么造的（重点）", "一个人 + Cursor 聊出来：人提需求 + 把关，AI 写几乎全部实现，小步迭代", WARN),
    ("最想分享的", "门槛降了 —— 你也能这样造内部工具；但方向、边界、把关，仍是人的活儿", ACCENT2),
]
y = int(1.95 * IN)
for t, d, ac in rows3:
    box(s, int(0.95 * IN), y, int(10.6 * IN), int(1.2 * IN), fill=BG_PANEL, line=LINE)
    box(s, int(0.95 * IN), y, int(0.12 * IN), int(1.2 * IN), fill=ac, radius=False)
    tf = textbox(s, int(1.35 * IN), y + int(0.16 * IN), int(10 * IN), int(0.95 * IN), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, t, size=19, color=ac, bold=True, first=True, space_after=4)
    para(tf, d, size=15, color=TEXT, space_after=0, line_spacing=1.05)
    y += int(1.45 * IN)

# ---- 16. 谢谢 ------------------------------------------------------------
s = add_slide(prs)
box(s, 0, 0, SLIDE_W, int(0.18 * IN), fill=ACCENT, radius=False)
box(s, 0, int(0.18 * IN), SLIDE_W, int(0.06 * IN), fill=ACCENT2, radius=False)
tf = textbox(s, int(0.9 * IN), int(2.6 * IN), int(11.0 * IN), int(2.0 * IN), anchor=MSO_ANCHOR.MIDDLE)
para(tf, "谢谢 · 欢迎提问", size=44, color=TEXT, bold=True, first=True, space_after=12)
para(tf, "现场演示 + Q&A", size=20, color=ACCENT, space_after=0)

# ---- 页脚 ----------------------------------------------------------------
total = n()
skip = {1, 7, total}  # 封面、重点分隔页、谢谢页
for i, slide in enumerate(prs.slides, start=1):
    if i in skip:
        continue
    footer(slide, i, total)

prs.save(str(OUT))
print(f"OK -> {OUT}  ({total} slides)")
